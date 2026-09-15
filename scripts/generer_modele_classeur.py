"""Genere un modele de classeur existant, a remplacer par le votre.

    python scripts/generer_modele_classeur.py

Produit templates/classeur/rapport_existant.xlsx avec la structure que le
script attend : une ligne d'en-tete en ligne 1, les donnees a partir de la
ligne 2, sur trois onglets.

Un quatrieme onglet contient des formules qui pointent vers les autres :
il sert a verifier que le collage ne casse rien autour de lui.

REMPLACEZ ce fichier par votre vrai classeur des que possible -- c'est lui
qui fait foi. Voir docs/ancienne-version-excel.md.
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RACINE = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RACINE / "src"))

from rapports.config import charger  # noqa: E402
from rapports.sorties import style  # noqa: E402

ONGLETS = {
    "Donnees": ["Date CAV", "Segment", "Code DR"],
    "Objectifs": ["Code DR", "Objectif debut annee", "Objectif annuel"],
    "Synthese": [
        "Code DR", "Realise semaine", "Realise cumul", "Objectif debut annee",
        "R/O", "Ecart", "Objectif annuel", "% objectif annuel",
    ],
}


def main() -> int:
    cfg = charger(RACINE)
    chemin = cfg.modele_classeur

    if chemin.exists():
        reponse = input(f"{chemin} existe deja. Ecraser ? [o/N] ").strip().lower()
        if reponse not in ("o", "oui", "y"):
            print("Abandon : le modele existant est conserve.")
            return 1

    classeur = Workbook()
    classeur.remove(classeur.active)

    for nom, entetes in ONGLETS.items():
        feuille = classeur.create_sheet(nom)
        for indice, entete in enumerate(entetes, start=1):
            cellule = feuille.cell(row=1, column=indice, value=entete)
            cellule.font = Font(bold=True, color="FFFFFF")
            cellule.fill = PatternFill("solid", fgColor=style.ROUGE.lstrip("#").upper())
            cellule.alignment = Alignment(horizontal="center")
            feuille.column_dimensions[get_column_letter(indice)].width = max(14, len(entete) + 3)
        feuille.freeze_panes = "A2"

    # Onglet temoin : des formules qui doivent survivre au collage.
    restitution = classeur.create_sheet("Restitution")
    restitution["A1"] = "Controle automatique"
    restitution["A1"].font = Font(bold=True, size=13)
    restitution["A3"] = "Lignes collees dans Donnees"
    restitution["B3"] = "=COUNTA(Donnees!C:C)-1"
    restitution["A4"] = "Nombre de DR"
    restitution["B4"] = "=COUNTA(Objectifs!A:A)-1"
    restitution["A5"] = "Realise cumul total"
    restitution["B5"] = "=SUM(Synthese!C:C)"
    restitution.column_dimensions["A"].width = 30
    restitution.column_dimensions["B"].width = 16

    chemin.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(chemin)
    print(f"Modele genere : {chemin}")
    print("Onglets :", ", ".join(classeur.sheetnames))
    print("\nRemplacez ce fichier par votre vrai classeur, en conservant les")
    print("noms d'onglets declares dans config/rapports.toml.")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
