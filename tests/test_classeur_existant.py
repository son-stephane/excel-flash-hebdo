"""Remplissage du classeur historique.

Ce qui compte ici n'est pas tant ce que le script ecrit que ce qu'il
n'abime pas : les en-tetes, les formules et les autres onglets du classeur
que l'on remplit aujourd'hui a la main.
"""

import datetime as dt

import pandas as pd
import pytest
from openpyxl import Workbook, load_workbook

from rapports import calculs
from rapports.config import OngletCollage
from rapports.sorties import classeur_existant
from fabrique import DANS_LA_SEMAINE, SEMAINE, ligne, objectifs, source

def _lignes_pleines(chemin, onglet: str) -> int:
    """Lignes de donnees reellement renseignees, en-tete exclu.

    On ne se fie pas a max_row : vider une cellule lui laisse son existence
    dans le fichier, donc la derniere ligne declaree ne diminue pas.
    """
    feuille = load_workbook(chemin)[onglet]
    return sum(
        1
        for ligne in feuille.iter_rows(min_row=2, values_only=True)
        if any(valeur is not None for valeur in ligne)
    )


ONGLETS = (
    OngletCollage("Donnees", "donnees", "activite", "A2"),
    OngletCollage("Objectifs", "objectifs", "activite", "A2"),
    OngletCollage("Synthese", "synthese", "activite", "A2"),
)


@pytest.fixture
def modele(tmp_path):
    """Un classeur avec en-tetes, plus un onglet de formules a preserver."""
    classeur = Workbook()
    classeur.remove(classeur.active)
    for nom, entetes in {
        "Donnees": ["Date CAV", "Segment", "Code DR"],
        "Objectifs": ["Code DR", "Objectif debut annee", "Objectif annuel"],
        "Synthese": ["Code DR", "Realise semaine", "Realise cumul", "Objectif debut annee",
                     "R/O", "Ecart", "Objectif annuel", "% objectif annuel"],
    }.items():
        feuille = classeur.create_sheet(nom)
        feuille.append(entetes)

    temoin = classeur.create_sheet("Restitution")
    temoin["A1"] = "Total"
    temoin["B1"] = "=SUM(Synthese!C:C)"

    chemin = tmp_path / "modele.xlsx"
    classeur.save(chemin)
    return chemin


@pytest.fixture
def resultat(rapport):
    return calculs.calculer(
        rapport,
        source([ligne(DANS_LA_SEMAINE, dr="DR01")] * 4 + [ligne(DANS_LA_SEMAINE, dr="DR02")] * 2),
        objectifs({"DR01": (3.0, 12.0), "DR02": (3.0, 12.0)}),
        SEMAINE,
    )


def remplir(resultat, modele, destination, onglets=ONGLETS):
    return classeur_existant.remplir(
        {"activite": resultat}, onglets, modele, destination, moteur="openpyxl"
    )


def test_les_trois_onglets_sont_remplis(resultat, modele, tmp_path):
    collage = remplir(resultat, modele, tmp_path / "sortie.xlsx")

    assert collage.onglets == ["Donnees", "Objectifs", "Synthese"]
    classeur = load_workbook(collage.chemin)
    assert classeur["Donnees"].max_row == 7        # 1 en-tete + 6 lignes
    assert classeur["Objectifs"].max_row == 3      # 1 en-tete + 2 DR
    # 1 en-tete + 2 DR + 1 total. Pas de ligne "Total hors DR99" : cette DR
    # n'existe pas dans le jeu de test, et le calcul l'a signale.
    assert classeur["Synthese"].max_row == 4


def test_la_ligne_d_entete_n_est_pas_ecrasee(resultat, modele, tmp_path):
    """Le collage commence en A2 : la ligne 1 du classeur reste celle d'origine."""
    collage = remplir(resultat, modele, tmp_path / "sortie.xlsx")
    feuille = load_workbook(collage.chemin)["Donnees"]

    assert [c.value for c in feuille[1]] == ["Date CAV", "Segment", "Code DR"]
    assert isinstance(feuille["A2"].value, dt.datetime)


def test_les_formules_des_autres_onglets_survivent(resultat, modele, tmp_path):
    collage = remplir(resultat, modele, tmp_path / "sortie.xlsx")
    temoin = load_workbook(collage.chemin)["Restitution"]

    assert temoin["B1"].value == "=SUM(Synthese!C:C)"


def test_le_modele_n_est_jamais_modifie(resultat, modele, tmp_path):
    avant = modele.read_bytes()
    remplir(resultat, modele, tmp_path / "sortie.xlsx")

    assert modele.read_bytes() == avant


def test_l_ancien_contenu_est_efface(resultat, modele, tmp_path):
    """Une semaine plus courte ne doit pas laisser trainer les lignes de la
    semaine precedente sous les nouvelles."""
    premiere = tmp_path / "semaine1.xlsx"
    remplir(resultat, modele, premiere)
    assert _lignes_pleines(premiere, "Donnees") == 6

    court = calculs.calculer(
        resultat.rapport,
        source([ligne(DANS_LA_SEMAINE, dr="DR01")]),
        objectifs({"DR01": (3.0, 12.0)}),
        SEMAINE,
    )
    # Le classeur deja rempli sert de modele : les anciennes lignes doivent
    # disparaitre, pas rester sous les nouvelles.
    seconde = tmp_path / "semaine2.xlsx"
    remplir(court, premiere, seconde)

    assert _lignes_pleines(seconde, "Donnees") == 1


def test_onglet_absent_signale_sans_interrompre(resultat, modele, tmp_path):
    onglets = ONGLETS + (OngletCollage("Inexistant", "donnees", "activite", "A2"),)
    collage = remplir(resultat, modele, tmp_path / "sortie.xlsx", onglets)

    assert "Inexistant" not in collage.onglets
    assert len(collage.onglets) == 3
    assert any("Inexistant" in message for message in collage.avertissements)


def test_modele_et_sortie_confondus_refuses(resultat, modele):
    """Ecrire dans le modele reviendrait a perdre le point de depart."""
    with pytest.raises(classeur_existant.ClasseurIndisponible, match="meme fichier"):
        remplir(resultat, modele, modele)


def test_modele_introuvable_dit_quoi_faire(resultat, tmp_path):
    with pytest.raises(classeur_existant.ClasseurIndisponible, match="generer_modele_classeur"):
        remplir(resultat, tmp_path / "absent.xlsx", tmp_path / "sortie.xlsx")


def test_bloc_synthese_sans_colonne_technique(resultat):
    synthese = classeur_existant.bloc(resultat, "synthese")
    assert "type_ligne" not in synthese.columns
    assert list(synthese.columns)[0] == "Code DR"


def test_contenu_inconnu_refuse(resultat):
    with pytest.raises(classeur_existant.ClasseurIndisponible, match="Contenu inconnu"):
        classeur_existant.bloc(resultat, "autre_chose")
