"""Bout en bout : des deux fichiers Excel aux trois sorties."""

import shutil
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from rapports import pipeline
from fabrique import AVANT_LA_SEMAINE, DANS_LA_SEMAINE, SEMAINE, ecrire_excel

GABARIT = Path(__file__).resolve().parents[1] / "templates" / "rapport.html.j2"


def _preparer(cfg):
    shutil.copy(GABARIT, cfg.racine / "templates" / GABARIT.name)
    ecrire_excel(
        cfg.dossier_entrees, "activite.xlsx",
        pd.DataFrame(
            {
                "Date CAV": [DANS_LA_SEMAINE] * 3 + [AVANT_LA_SEMAINE] * 7,
                "Segment": ["Grand Public"] * 10,
                "Code DR": ["DR01"] * 6 + ["DR99"] * 4,
            }
        ),
    )
    ecrire_excel(
        cfg.dossier_entrees, "objectifs.xlsx",
        pd.DataFrame(
            {
                "Code DR": ["DR01", "DR99"],
                "Objectif debut annee": [5, 5],
                "Objectif annuel": [20, 20],
            }
        ),
    )


def test_les_trois_sorties_sont_produites(cfg):
    _preparer(cfg)
    sortie = pipeline.executer(cfg, "activite", SEMAINE)

    assert sortie.classeur.exists()
    assert sortie.page_html.exists()
    assert sortie.image.exists()
    assert sortie.classeur.parent.name == SEMAINE  # sorties rangees par semaine


def test_le_classeur_contient_le_tableau_et_un_graphique(cfg):
    _preparer(cfg)
    sortie = pipeline.executer(cfg, "activite", SEMAINE)

    feuille = load_workbook(sortie.classeur)["Synthese"]
    assert len(feuille._charts) == 1

    entetes = [c.value for c in feuille[5]]
    assert entetes[0] == "Code DR" and "R/O" in entetes

    codes = [feuille.cell(row=r, column=1).value for r in range(6, 6 + 4)]
    assert codes == ["DR01", "DR99", "Total toutes DR", "Total hors DR99"]


def test_la_page_html_est_autonome(cfg):
    """Aucune ressource externe : le fichier doit rester lisible hors ligne."""
    _preparer(cfg)
    sortie = pipeline.executer(cfg, "activite", SEMAINE)
    page = sortie.page_html.read_text(encoding="utf-8")

    assert "data:image/png;base64," in page
    assert "http://" not in page and "https://" not in page
    assert "Total hors DR99" in page


def test_relancer_ecrase_sans_dupliquer(cfg):
    _preparer(cfg)
    pipeline.executer(cfg, "activite", SEMAINE)
    pipeline.executer(cfg, "activite", SEMAINE)

    produits = list((cfg.dossier_sorties / SEMAINE).glob("activite_*"))
    assert len(produits) == 3
