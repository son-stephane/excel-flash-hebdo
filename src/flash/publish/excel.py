"""Publication vers le classeur Excel (phase 1).

Deux etapes volontairement separees :

  * INJECTION (openpyxl)  -- ecrit le mart dans l'onglet plat et redimensionne
    le tableau structure. Ne demande pas Excel : marche sur un serveur.
  * RAFRAICHISSEMENT (xlwings) -- ouvre le classeur, rafraichit les TCD et
    exporte les graphiques en PNG. Demande Excel installe.

Si Excel n'est pas disponible, l'injection est quand meme faite et le
pipeline bascule sur les graphiques Python : on ne perd pas la semaine.

PREREQUIS DU TEMPLATE : l'onglet de donnees doit contenir un TABLEAU
STRUCTURE (Insertion > Tableau) nomme comme dans settings.toml, et les TCD
doivent pointer sur CE TABLEAU -- pas sur une plage figee du type A1:M5000,
sinon ils ne suivront pas la variation du nombre de lignes.
"""

from __future__ import annotations

import shutil
from dataclasses import dataclass, field
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from ..config import Config
from ..logging_setup import journal
from ..paths import Emplacements
from ..transform import marts

LOG = journal("excel")


class ExcelIndisponible(RuntimeError):
    """Excel n'est pas pilotable sur cette machine."""


class TemplateInvalide(RuntimeError):
    pass


@dataclass
class ResultatExcel:
    classeur: Path
    lignes_injectees: int
    tcd_rafraichis: bool = False
    graphiques: list[Path] = field(default_factory=list)


# --------------------------------------------------------------------------
def preparer_classeur(cfg: Config, semaine: str) -> Path:
    """Copie le template vers le classeur de la semaine. Le template n'est jamais modifie."""
    lieux = Emplacements(cfg)
    template = lieux.template_excel()
    if not template.exists():
        raise TemplateInvalide(
            f"Template introuvable : {template}\n"
            f"  Le generer avec : python scripts/make_template.py"
        )
    destination = Emplacements.preparer(lieux.classeur(semaine))
    shutil.copy2(template, destination)
    LOG.debug("Classeur de la semaine : %s", destination)
    return destination


def injecter(cfg: Config, classeur: Path, donnees: pd.DataFrame) -> int:
    """Ecrit le mart dans l'onglet plat et redimensionne le tableau structure."""
    reglages = cfg.bloc("excel")
    nom_onglet = reglages.get("onglet_donnees", "DONNEES")
    nom_tableau = reglages.get("tableau_donnees", "tbl_donnees")

    wb = load_workbook(classeur)
    if nom_onglet not in wb.sheetnames:
        raise TemplateInvalide(
            f"Onglet '{nom_onglet}' absent du classeur. Onglets presents : {wb.sheetnames}"
        )
    feuille = wb[nom_onglet]

    # On vide l'ancienne zone de donnees en gardant la ligne d'en-tete.
    if feuille.max_row > 1:
        feuille.delete_rows(2, feuille.max_row - 1)

    entetes = [cellule.value for cellule in feuille[1]]
    if not any(entetes):
        entetes = list(donnees.columns)
        for indice, nom in enumerate(entetes, start=1):
            feuille.cell(row=1, column=indice, value=nom)
    else:
        manquantes = [c for c in entetes if c and c not in donnees.columns]
        if manquantes:
            raise TemplateInvalide(
                f"Colonnes attendues par le template absentes du mart : {manquantes}\n"
                f"  Colonnes du mart : {list(donnees.columns)}\n"
                f"  -> aligner transform/sql/mart_flash_hebdo.sql et le template."
            )
        entetes = [c for c in entetes if c]

    ordonne = donnees[entetes]
    for indice_ligne, ligne in enumerate(ordonne.itertuples(index=False), start=2):
        for indice_colonne, valeur in enumerate(ligne, start=1):
            feuille.cell(row=indice_ligne, column=indice_colonne, value=_valeur_excel(valeur))

    _redimensionner_tableau(feuille, nom_tableau, len(ordonne), len(entetes))
    wb.save(classeur)
    LOG.info("%d lignes injectees dans '%s'", len(ordonne), nom_onglet)
    return len(ordonne)


def _valeur_excel(valeur):
    if valeur is None or (isinstance(valeur, float) and pd.isna(valeur)):
        return None
    if isinstance(valeur, pd.Timestamp):
        return valeur.to_pydatetime()
    if hasattr(valeur, "item"):  # types numpy
        return valeur.item()
    return valeur


def _redimensionner_tableau(feuille, nom_tableau: str, lignes: int, colonnes: int) -> None:
    tableaux = getattr(feuille, "tables", {})
    if nom_tableau not in tableaux:
        LOG.warning(
            "Tableau structure '%s' absent de l'onglet : les TCD risquent de ne pas "
            "suivre le nombre de lignes. Voir scripts/make_template.py",
            nom_tableau,
        )
        return
    derniere = max(lignes + 1, 2)
    tableaux[nom_tableau].ref = f"A1:{get_column_letter(colonnes)}{derniere}"
    LOG.debug("Tableau '%s' redimensionne sur %s", nom_tableau, tableaux[nom_tableau].ref)


# --------------------------------------------------------------------------
def rafraichir_et_exporter(cfg: Config, classeur: Path, dossier_png: Path) -> list[Path]:
    """Rafraichit les TCD puis exporte les graphiques du classeur en PNG."""
    try:
        import xlwings as xw
    except ImportError as exc:
        raise ExcelIndisponible(
            "xlwings n'est pas installe. Sur le poste Windows : pip install .[windows]"
        ) from exc

    reglages = cfg.bloc("excel")
    voulus = set(reglages.get("graphiques_exportes") or [])
    dossier_png.mkdir(parents=True, exist_ok=True)
    exportes: list[Path] = []

    app = None
    livre = None
    try:
        app = xw.App(visible=bool(reglages.get("visible", False)), add_book=False)
        app.display_alerts = False
        app.screen_updating = False
        livre = app.books.open(str(classeur))

        livre.api.RefreshAll() if hasattr(livre, "api") else None
        app.api.CalculateUntilAsyncQueriesDone() if hasattr(app, "api") else None
        LOG.info("TCD et connexions rafraichis")

        for feuille in livre.sheets:
            for graphique in feuille.charts:
                if voulus and graphique.name not in voulus:
                    continue
                cible = dossier_png / f"{_nom_fichier(feuille.name, graphique.name)}.png"
                _exporter_graphique(graphique, cible)
                exportes.append(cible)
                LOG.debug("Graphique exporte : %s", cible.name)

        livre.save()
    except Exception as exc:  # remonte avec un message actionnable
        raise ExcelIndisponible(
            f"Echec du pilotage d'Excel : {exc}\n"
            f"  Verifier qu'Excel est installe, qu'aucune boite de dialogue n'est "
            f"ouverte, et que le classeur n'est pas deja ouvert."
        ) from exc
    finally:
        if livre is not None:
            livre.close()
        if app is not None:
            app.quit()

    LOG.info("%d graphique(s) exportes depuis le classeur", len(exportes))
    return exportes


def _exporter_graphique(graphique, cible: Path) -> None:
    """L'API d'export differe entre Windows et macOS."""
    api = graphique.api
    for tentative in (
        lambda: api[1].Export(str(cible)),   # Windows : (ChartObject, Chart)
        lambda: api.Export(str(cible)),      # Windows, variante
        lambda: api.export(str(cible)),      # macOS / appscript
    ):
        try:
            tentative()
            return
        except (AttributeError, TypeError, IndexError):
            continue
    raise ExcelIndisponible(
        f"Impossible d'exporter le graphique '{graphique.name}'. "
        f"L'export de graphiques est fiable sur Windows ; sur macOS, utiliser "
        f"le moteur 'python' (settings.toml > publication.moteur_graphiques)."
    )


def _nom_fichier(feuille: str, graphique: str) -> str:
    brut = f"{feuille}_{graphique}"
    return "".join(c if c.isalnum() or c in "-_" else "_" for c in brut)


# --------------------------------------------------------------------------
def publier(cfg: Config, semaine: str, *, avec_excel: bool = True) -> ResultatExcel:
    """Chaine complete : copie du template, injection, rafraichissement, export."""
    donnees = marts.lire_mart(cfg, semaine)
    classeur = preparer_classeur(cfg, semaine)
    lignes = injecter(cfg, classeur, donnees)
    resultat = ResultatExcel(classeur=classeur, lignes_injectees=lignes)

    if not avec_excel:
        return resultat

    dossier_png = Emplacements(cfg).graphiques(semaine)
    try:
        resultat.graphiques = rafraichir_et_exporter(cfg, classeur, dossier_png)
        resultat.tcd_rafraichis = True
    except ExcelIndisponible as exc:
        LOG.warning("Excel indisponible : %s", exc)
        LOG.warning("Bascule sur les graphiques Python (moteur de secours).")
        from . import charts

        resultat.graphiques = charts.produire(cfg, semaine)

    return resultat
