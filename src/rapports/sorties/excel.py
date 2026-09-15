"""Classeur Excel du rapport : tableau par DR et graphique du R/O.

Le graphique est un VRAI graphique Excel, pas une image : il reste lie aux
cellules, donc modifiable et restylable par le destinataire.
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..calculs import LIGNE_DR, Resultat
from ..journal import journal
from ..semaines import libelle as libelle_semaine
from . import style

LOG = journal("excel")

ONGLET = "Synthese"

COLONNES = [
    ("code_dr", "Code DR", None, 16),
    ("realise_semaine", "Realise semaine", style.FORMAT_ENTIER, 16),
    ("realise_cumul", "Realise cumul", style.FORMAT_ENTIER, 15),
    ("objectif_debut_annee", "Objectif debut annee", style.FORMAT_ENTIER, 20),
    ("ro", "R/O", style.FORMAT_POURCENT, 10),
    ("ecart", "Ecart", style.FORMAT_ECART, 12),
    ("objectif_annuel", "Objectif annuel", style.FORMAT_ENTIER, 16),
    ("part_annuel", "% objectif annuel", style.FORMAT_POURCENT, 17),
]

LIGNE_ENTETE = 5


def _sans_diese(couleur: str) -> str:
    return couleur.lstrip("#").upper()


def produire(resultat: Resultat, chemin: Path) -> Path:
    classeur = Workbook()
    feuille = classeur.active
    feuille.title = ONGLET

    _entete_document(feuille, resultat)
    _tableau(feuille, resultat)
    _graphique(feuille, resultat)

    feuille.freeze_panes = f"A{LIGNE_ENTETE + 1}"
    chemin.parent.mkdir(parents=True, exist_ok=True)
    classeur.save(chemin)
    LOG.info("Classeur ecrit : %s", chemin)
    return chemin


def _entete_document(feuille, resultat: Resultat) -> None:
    rapport = resultat.rapport

    feuille["A1"] = rapport.libelle
    feuille["A1"].font = Font(bold=True, size=15, color=_sans_diese(style.ENCRE))

    feuille["A2"] = f"Semaine {resultat.semaine} - {libelle_semaine(resultat.semaine)}"
    feuille["A2"].font = Font(size=11, color=_sans_diese(style.ENCRE_DOUCE))

    segments = ", ".join(rapport.segments) if rapport.segments else "tous"
    feuille["A3"] = (
        f"Segments retenus : {segments}"
        f"   |   Cumul du 1er janvier au {libelle_semaine(resultat.semaine).split('au ')[-1]}"
        f"   |   {resultat.lignes_apres_filtre} ligne(s) retenues sur {resultat.lignes_lues}"
    )
    feuille["A3"].font = Font(size=9, italic=True, color=_sans_diese(style.ENCRE_PALE))


def _tableau(feuille, resultat: Resultat) -> None:
    fond_entete = PatternFill("solid", fgColor=_sans_diese(style.BLEU))
    fond_total = PatternFill("solid", fgColor=_sans_diese(style.FOND_DOUX))
    filet = Side(style="thin", color=_sans_diese(style.TRAIT))
    bordure = Border(bottom=filet)

    for indice, (_, entete, _, largeur) in enumerate(COLONNES, start=1):
        cellule = feuille.cell(row=LIGNE_ENTETE, column=indice, value=entete)
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = fond_entete
        cellule.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        feuille.column_dimensions[get_column_letter(indice)].width = largeur
    feuille.row_dimensions[LIGNE_ENTETE].height = 30

    for decalage, (_, ligne) in enumerate(resultat.tableau.iterrows()):
        numero = LIGNE_ENTETE + 1 + decalage
        est_total = ligne["type_ligne"] != LIGNE_DR

        for indice, (champ, _, format_nombre, _) in enumerate(COLONNES, start=1):
            valeur = ligne[champ]
            cellule = feuille.cell(
                row=numero, column=indice,
                value=None if pd.isna(valeur) else _valeur(valeur),
            )
            if format_nombre:
                cellule.number_format = format_nombre
            cellule.border = bordure
            if est_total:
                cellule.font = Font(bold=True)
                cellule.fill = fond_total
            if champ == "ecart" and not pd.isna(valeur):
                cellule.font = Font(
                    bold=est_total, color=_sans_diese(style.couleur_ecart(valeur))
                )
        feuille.cell(row=numero, column=1).alignment = Alignment(horizontal="left")


def _valeur(valeur):
    """Convertit les types numpy en types Python, qu'openpyxl sait ecrire."""
    return valeur.item() if hasattr(valeur, "item") else valeur


def _graphique(feuille, resultat: Resultat) -> None:
    nombre_dr = len(resultat.par_dr)
    if nombre_dr == 0:
        return

    premiere = LIGNE_ENTETE + 1
    derniere = LIGNE_ENTETE + nombre_dr  # les totaux sont exclus du graphique
    colonne_ro = next(i for i, c in enumerate(COLONNES, start=1) if c[0] == "ro")

    histogramme = BarChart()
    histogramme.type = "col"
    histogramme.title = "Realise / objectif par DR"
    histogramme.height = 9
    histogramme.width = max(16, 2.2 * nombre_dr)
    histogramme.legend = None  # une seule serie : la legende n'apprend rien
    histogramme.y_axis.numFmt = style.FORMAT_POURCENT
    histogramme.y_axis.majorGridlines = None

    serie = Series(
        Reference(feuille, min_col=colonne_ro, min_row=premiere, max_row=derniere),
        title="R/O",
    )
    histogramme.append(serie)
    histogramme.set_categories(
        Reference(feuille, min_col=1, min_row=premiere, max_row=derniere)
    )

    feuille.add_chart(histogramme, f"A{derniere + 4}")
